"""Vistas GESAC Campus: importador de 'Detalle de Pagos', cuadre contra Fiserv,
reportes de ingresos (forma de pago / categoria) y cartera vencida detallada."""

import datetime
from collections import OrderedDict, defaultdict
from decimal import Decimal

import xlrd
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Count, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font

from conciliaciones.models import SaldoCuentaPeriodo
from conciliaciones.utils import validar_periodo_abierto
from empresas.models import CuentaBancaria

from . import ingresos_utils as iu
from .decoradores import requiere_membresia_activa_escuela
from .models import Alumno, CarteraVencida, IngresoEscuela, LoteImportacionIngresos

CERO = Decimal('0.00')


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def _empresa_escuela(request):
    perfil = getattr(request.user, 'perfilusuario', None)
    if not perfil or not perfil.empresa or perfil.empresa.segmento != 'escuela':
        return None
    return perfil.empresa


def _fecha_get(request, nombre, defecto):
    raw = request.GET.get(nombre)
    if raw:
        try:
            return datetime.date.fromisoformat(raw)
        except ValueError:
            pass
    return defecto


def _excel_response(wb, nombre):
    resp = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    resp['Content-Disposition'] = f'attachment; filename={nombre}'
    wb.save(resp)
    return resp


def _encabezado(ws, columnas):
    ws.append(columnas)
    for c in ws[1]:
        c.font = Font(bold=True)


# ----------------------------------------------------------------------------
# 1) Importar Detalle de Pagos (IPC / Cafeteria / Colegiaturas)
# ----------------------------------------------------------------------------
@login_required
@requiere_membresia_activa_escuela
def importar_ingresos_excel(request):
    empresa = _empresa_escuela(request)
    if not empresa:
        messages.error(request, 'No tienes una empresa de escuela asociada.')
        return redirect('dashboard_inicio')

    cuentas = CuentaBancaria.objects.filter(empresa=empresa, activa=True)

    if request.method == 'POST':
        archivo = request.FILES.get('archivo_excel')
        fuente = request.POST.get('fuente', '')
        cuenta_id = request.POST.get('cuenta_bancaria', '')

        if not archivo:
            messages.error(request, 'Selecciona un archivo .xls.')
            return redirect('importar_ingresos_excel')
        if fuente not in dict(IngresoEscuela.FUENTE_CHOICES):
            messages.error(request, 'Elige la fuente del archivo (IPC, Cafetería o Colegiaturas).')
            return redirect('importar_ingresos_excel')

        try:
            wb = xlrd.open_workbook(file_contents=archivo.read())
            parseo = iu.parsear_detalle_pagos(wb)
        except Exception as e:  # noqa: BLE001
            messages.error(request, f'No se pudo leer el archivo -- verifica que sea el "Detalle de Pagos" .xls de Academic+ ({e}).')
            return redirect('importar_ingresos_excel')

        if parseo['descuadres']:
            messages.error(
                request,
                'El archivo no se importó porque los subtotales no cuadran con las filas leídas: '
                + ' | '.join(parseo['descuadres']),
            )
            return redirect('importar_ingresos_excel')
        if parseo['errores']:
            messages.error(
                request,
                'No se importó nada porque hay filas que no se pudieron leer (corrígelas en el archivo): '
                + ' | '.join(parseo['errores'][:8]),
            )
            return redirect('importar_ingresos_excel')
        if not parseo['filas']:
            messages.warning(request, 'El archivo no tenía cobros para importar.')
            return redirect('importar_ingresos_excel')

        # Cuenta por defecto (opcional): solo se usa si la columna "Cuenta Depósito"
        # del archivo no existe o no coincide con una cuenta registrada.
        cuenta_defecto = cuentas.filter(pk=cuenta_id).first() if cuenta_id.isdigit() else None
        errores_cuenta = iu.resolver_cuentas(parseo, cuentas, cuenta_defecto)
        if errores_cuenta:
            messages.error(request, 'No se importó nada. ' + ' | '.join(errores_cuenta))
            return redirect('importar_ingresos_excel')
        cuenta = cuenta_defecto

        # Reglas de fechas del sistema (no futuras, no anios anteriores).
        pares = {(f['cuenta'], f['fecha']) for f in parseo['filas']}
        for cta, fecha in sorted(pares, key=lambda p: (p[1], p[0].pk if p[0] else 0)):
            ok, error = validar_periodo_abierto(cta, fecha)
            if not ok:
                messages.error(request, f"{error} (Fecha Abono {fecha.strftime('%d/%m/%Y')}). No se importó nada.")
                return redirect('importar_ingresos_excel')

        # Aviso (no bloquea) si algun cobro cae en un mes con periodo bancario cerrado.
        avisos_cerrado = []
        meses = {
            (f['cuenta'], f['fecha'].year, f['fecha'].month)
            for f in parseo['filas'] if f['cuenta'] is not None
        }
        for cta, anio, mes in sorted(meses, key=lambda m: (m[1], m[2], m[0].pk)):
            p = SaldoCuentaPeriodo.objects.filter(cuenta=cta, anio=anio, mes=mes, cerrado=True).first()
            if p:
                avisos_cerrado.append(f'{p.nombre_mes()} {anio} ({cta})')

        try:
            with transaction.atomic():
                r = iu.importar_filas(
                    empresa=empresa, usuario=request.user, fuente=fuente,
                    archivo_nombre=archivo.name, cuenta_bancaria=cuenta, parseo=parseo,
                )
        except Exception as e:  # noqa: BLE001
            messages.error(request, f'No se importó nada por un error: {e}')
            return redirect('importar_ingresos_excel')

        messages.success(
            request,
            f"Se importaron {r['creadas']} cobro(s): ${r['total_efectivo']:,.2f} en efectivo → caja chica "
            f"y ${r['total_banco']:,.2f} a bancos (100%, sin descontar comisión).",
        )
        if r['duplicadas']:
            messages.info(request, f"{r['duplicadas']} fila(s) ya estaban importadas y se omitieron (sin duplicar).")
        if parseo['omitidas']:
            messages.info(request, f"{len(parseo['omitidas'])} fila(s) con importe cero o negativo no se importaron.")
        if parseo['errores']:
            messages.warning(request, ' | '.join(parseo['errores'][:8]))
        if avisos_cerrado:
            messages.warning(request, 'Ojo: hay cobros con Fecha Abono en periodo(s) bancario(s) ya cerrado(s): ' + ', '.join(avisos_cerrado) + '.')
        if r['sin_clasificar']:
            messages.warning(request, 'Conceptos sin categoría reconocida (quedaron como "Sin clasificar"): ' + ' | '.join(r['sin_clasificar'][:8]))
        return redirect('importar_ingresos_excel')

    lotes = LoteImportacionIngresos.objects.filter(empresa=empresa).select_related('cuenta_bancaria')[:15]
    return render(request, 'escuelas/importar_ingresos.html', {
        'cuentas': cuentas,
        'fuentes': IngresoEscuela.FUENTE_CHOICES,
        'lotes': lotes,
    })


@login_required
@requiere_membresia_activa_escuela
@require_POST
def deshacer_lote_ingresos(request, lote_id):
    empresa = _empresa_escuela(request)
    if not empresa:
        return redirect('dashboard_inicio')
    lote = get_object_or_404(LoteImportacionIngresos, pk=lote_id, empresa=empresa)
    try:
        with transaction.atomic():
            nombre = lote.archivo
            iu.deshacer_lote(lote)
        messages.success(request, f'Se deshizo la importación de "{nombre}".')
    except ValueError as e:
        messages.error(request, str(e))
    return redirect('importar_ingresos_excel')


# ----------------------------------------------------------------------------
# 2) Cuadre contra Fiserv (NO crea ingresos)
# ----------------------------------------------------------------------------
@login_required
@requiere_membresia_activa_escuela
def importar_ingresos_fiserv(request):
    empresa = _empresa_escuela(request)
    if not empresa:
        messages.error(request, 'No tienes una empresa de escuela asociada.')
        return redirect('dashboard_inicio')

    contexto = {'comisiones': iu.COMISION_REFERENCIA}
    if request.method == 'POST':
        archivo = request.FILES.get('archivo_excel')
        if not archivo:
            messages.error(request, 'Selecciona un archivo .xls.')
            return redirect('importar_ingresos_fiserv')
        try:
            wb = xlrd.open_workbook(file_contents=archivo.read())
            hoja = wb.sheet_by_index(0)
            encabezados = {iu.norm(hoja.cell(0, c).value): c for c in range(hoja.ncols)}
            requeridos = ['PAYW_RESULT', 'AMOUNT', 'NORECIBO', 'MERCHANT_ID', 'CUST_REQ_DATE']
            faltan = [r for r in requeridos if r not in encabezados]
            if faltan:
                raise ValueError('faltan columnas: ' + ', '.join(faltan))
        except Exception as e:  # noqa: BLE001
            messages.error(request, f'No se pudo leer el archivo de la plataforma ({e}).')
            return redirect('importar_ingresos_fiserv')

        def val(i, col):
            idx = encabezados.get(col)
            return hoja.cell(i, idx).value if idx is not None else ''

        aprobados, omitidos = [], 0
        for i in range(1, hoja.nrows):
            alumno = str(val(i, 'ALUMNO') or '')
            if 'PRUEBA' in iu.norm(alumno):
                continue
            if iu.norm(val(i, 'PAYW_RESULT')) != 'APROBADO':
                omitidos += 1
                continue
            monto = iu._dec(val(i, 'AMOUNT'))
            if not monto or monto <= 0:
                continue
            try:
                fecha = xlrd.xldate_as_datetime(val(i, 'CUST_REQ_DATE'), wb.datemode).date()
            except (ValueError, TypeError):
                fecha = None
            recibo = (str(val(i, 'NORECIBO')).split() or [''])[0]
            aprobados.append({
                'fecha': fecha, 'recibo': recibo, 'monto': monto, 'alumno': alumno,
                'matricula': iu._texto_id(val(i, 'AMATRICULA')),
                'descripcion': str(val(i, 'DESCRIPCIONPAGO') or ''),
                'merchant': str(val(i, 'MERCHANT_ID') or ''),
                'tipo': iu.tipo_tarjeta_merchant(val(i, 'MERCHANT_ID')),
            })

        # Resumen por tipo de tarjeta + comision de REFERENCIA (se registra a mano como gasto).
        por_tipo = OrderedDict((t, {'n': 0, 'total': CERO}) for t in ('debito', 'credito', 'amex'))
        por_tipo['otro'] = {'n': 0, 'total': CERO}
        for a in aprobados:
            g = por_tipo[a['tipo'] or 'otro']
            g['n'] += 1
            g['total'] += a['monto']
        resumen_tipos, comision_total = [], CERO
        etiquetas = {'debito': 'Débito', 'credito': 'Crédito', 'amex': 'AMEX', 'otro': 'Sin identificar'}
        for tipo, g in por_tipo.items():
            pct = iu.COMISION_REFERENCIA.get(tipo)
            com = (g['total'] * pct / Decimal('100')).quantize(Decimal('0.01')) if pct else CERO
            comision_total += com
            if g['n']:
                resumen_tipos.append({'etiqueta': etiquetas[tipo], 'n': g['n'], 'total': g['total'], 'pct': pct, 'comision': com})

        # Cuadre contra lo ya importado del Detalle de Pagos (cobros "portal web").
        web = IngresoEscuela.objects.filter(
            empresa=empresa, recibo__isnull=False, forma_pago__icontains='portal web'
        ).exclude(origen='fiserv')
        por_recibo_sistema = defaultdict(lambda: CERO)
        for r, m in web.values_list('recibo', 'monto'):
            por_recibo_sistema[r] += m
        todos_recibos = defaultdict(lambda: CERO)
        for r, m in IngresoEscuela.objects.filter(
            empresa=empresa, recibo__in=[a['recibo'] for a in aprobados]
        ).exclude(origen='fiserv').values_list('recibo', 'monto'):
            todos_recibos[r] += m

        sin_cobranza, monto_distinto, ok = [], [], 0
        for a in aprobados:
            if a['recibo'] not in todos_recibos:
                sin_cobranza.append(a)
            elif todos_recibos[a['recibo']] != a['monto']:
                monto_distinto.append({**a, 'monto_sistema': todos_recibos[a['recibo']]})
            else:
                ok += 1

        fechas = [a['fecha'] for a in aprobados if a['fecha']]
        sin_fiserv = []
        if fechas:
            recibos_fiserv = {a['recibo'] for a in aprobados}
            rango = web.filter(fecha__gte=min(fechas) - datetime.timedelta(days=1), fecha__lte=max(fechas) + datetime.timedelta(days=1))
            vistos = set()
            for r, alumno, fecha in rango.values_list('recibo', 'alumno', 'fecha'):
                if r not in recibos_fiserv and r not in vistos:
                    vistos.add(r)
                    sin_fiserv.append({'recibo': r, 'alumno': alumno, 'fecha': fecha, 'monto': por_recibo_sistema[r]})

        contexto.update({
            'resultado': True, 'archivo': archivo.name,
            'aprobados': len(aprobados), 'omitidos': omitidos,
            'total_aprobado': sum((a['monto'] for a in aprobados), CERO),
            'resumen_tipos': resumen_tipos, 'comision_total': comision_total,
            'ok': ok, 'sin_cobranza': sin_cobranza, 'monto_distinto': monto_distinto,
            'sin_fiserv': sin_fiserv,
        })
    return render(request, 'escuelas/importar_ingresos_fiserv.html', contexto)


# ----------------------------------------------------------------------------
# 3) Reportes de ingresos: por forma de pago / por categoria
# ----------------------------------------------------------------------------
def _atajos_periodo(hoy):
    ini_mes = hoy.replace(day=1)
    fin_mes_ant = ini_mes - datetime.timedelta(days=1)
    return [
        ('Mes actual', ini_mes, hoy),
        ('Mes anterior', fin_mes_ant.replace(day=1), fin_mes_ant),
        ('Año en curso', hoy.replace(month=1, day=1), hoy),
    ]


def _datos_ingresos(empresa, fecha_inicio, fecha_fin, fuente):
    # El origen 'fiserv' (importador anterior) duplica pagos que ya vienen en
    # el Detalle de Pagos, por eso no entra a estos reportes.
    qs = IngresoEscuela.objects.filter(
        empresa=empresa, fecha__gte=fecha_inicio, fecha__lte=fecha_fin
    ).exclude(origen='fiserv')
    if fuente:
        qs = qs.filter(fuente=fuente)
    filas = qs.values('forma_pago', 'rubro', 'concepto', 'es_efectivo').annotate(total=Sum('monto'), n=Count('id'))
    return qs, filas


def _reporte_ingresos(request, dim):
    """dim = 'forma_pago' (filas = forma de pago, columnas = categoria) o
    'rubro' (filas = categoria, columnas = forma de pago)."""
    empresa = _empresa_escuela(request)
    if not empresa:
        return redirect('dashboard_inicio')

    hoy = datetime.date.today()  # noqa: DTZ011
    fecha_inicio = _fecha_get(request, 'fecha_inicio', hoy.replace(day=1))
    fecha_fin = _fecha_get(request, 'fecha_fin', hoy)
    fuente = request.GET.get('fuente', '')
    if fuente not in dict(IngresoEscuela.FUENTE_CHOICES):
        fuente = ''

    qs, filas = _datos_ingresos(empresa, fecha_inicio, fecha_fin, fuente)

    celdas = defaultdict(lambda: CERO)
    n_mov = defaultdict(int)
    tot_fila, tot_col = defaultdict(lambda: CERO), defaultdict(lambda: CERO)
    destino_fila = {}
    for f in filas:
        forma = iu.etiqueta_forma_pago(f['forma_pago'])
        rubro = f['rubro'] or iu.normalizar_rubro(f['concepto'])
        fila_k, col_k = (forma, rubro) if dim == 'forma_pago' else (rubro, forma)
        celdas[(fila_k, col_k)] += f['total']
        n_mov[fila_k] += f['n']
        tot_fila[fila_k] += f['total']
        tot_col[col_k] += f['total']
        if dim == 'forma_pago':
            destino_fila[fila_k] = 'Caja chica' if f['es_efectivo'] else 'Bancos'

    gran_total = sum(tot_fila.values(), CERO)
    columnas = sorted(tot_col, key=lambda k: -tot_col[k])
    filas_out = []
    for k in sorted(tot_fila, key=lambda k: -tot_fila[k]):
        filas_out.append({
            'etiqueta': k, 'n': n_mov[k], 'total': tot_fila[k],
            'pct': (tot_fila[k] / gran_total * 100) if gran_total else CERO,
            'destino': destino_fila.get(k, ''),
            'celdas': [celdas.get((k, c), CERO) for c in columnas],
        })

    # Resumen de destino: efectivo (caja chica, por fuente) y bancos (por cuenta)
    total_efectivo = qs.filter(es_efectivo=True).aggregate(t=Sum('monto'))['t'] or CERO
    efectivo_fuentes = list(
        qs.filter(es_efectivo=True).values('fuente').annotate(t=Sum('monto')).order_by('-t')
    )
    etiquetas_fuente = dict(IngresoEscuela.FUENTE_CHOICES)
    for e in efectivo_fuentes:
        e['etiqueta'] = etiquetas_fuente.get(e['fuente'], 'Sin fuente')
    cuentas_map = {c.pk: str(c) for c in CuentaBancaria.objects.filter(empresa=empresa)}
    bancos = [
        {'cuenta': cuentas_map.get(b['cuenta_bancaria_id'], 'Sin cuenta'), 't': b['t']}
        for b in qs.filter(es_efectivo=False).values('cuenta_bancaria_id').annotate(t=Sum('monto')).order_by('-t')
    ]
    total_bancos = sum((b['t'] for b in bancos), CERO)

    titulo = 'Ingresos por forma de pago' if dim == 'forma_pago' else 'Ingresos por categoría'
    col_titulo = 'Categoría' if dim == 'forma_pago' else 'Forma de pago'
    fila_titulo = 'Forma de pago' if dim == 'forma_pago' else 'Categoría'

    if request.GET.get('exportar') == 'xlsx':
        wb = Workbook()
        ws = wb.active
        ws.title = titulo[:31]
        ws.append([titulo, f"{fecha_inicio:%d/%m/%Y} al {fecha_fin:%d/%m/%Y}"])
        _encabezado(ws, [fila_titulo, '# cobros', *columnas, 'Total', '% del total'] + (['Destino'] if dim == 'forma_pago' else []))
        for r in filas_out:
            ws.append([r['etiqueta'], r['n'], *[float(c) for c in r['celdas']], float(r['total']),
                       float(r['pct']) / 100] + ([r['destino']] if dim == 'forma_pago' else []))
        ws.append(['TOTAL', sum(r['n'] for r in filas_out), *[float(tot_col[c]) for c in columnas], float(gran_total), 1])
        ws.append([])
        ws.append(['Efectivo → caja chica', float(total_efectivo)])
        ws.append(['Bancos', float(total_bancos)])
        return _excel_response(wb, f"ingresos_{'forma_pago' if dim == 'forma_pago' else 'categoria'}.xlsx")

    return render(request, 'escuelas/reporte_ingresos.html', {
        'titulo': titulo, 'fila_titulo': fila_titulo, 'col_titulo': col_titulo,
        'dim': dim, 'filas': filas_out, 'columnas': columnas,
        'totales_col': [tot_col[c] for c in columnas], 'gran_total': gran_total,
        'total_movs': sum(r['n'] for r in filas_out),
        'total_efectivo': total_efectivo, 'efectivo_fuentes': efectivo_fuentes,
        'bancos': bancos, 'total_bancos': total_bancos,
        'fecha_inicio': fecha_inicio.isoformat(), 'fecha_fin': fecha_fin.isoformat(),
        'fuente': fuente, 'fuentes': IngresoEscuela.FUENTE_CHOICES,
        'atajos': _atajos_periodo(hoy),
        'labels_grafico': [r['etiqueta'] for r in filas_out],
        'datos_grafico': [float(r['total']) for r in filas_out],
    })


@login_required
@requiere_membresia_activa_escuela
def reporte_ingresos_forma_pago(request):
    return _reporte_ingresos(request, 'forma_pago')


@login_required
@requiere_membresia_activa_escuela
def reporte_ingresos_categoria(request):
    return _reporte_ingresos(request, 'rubro')


# ----------------------------------------------------------------------------
# 4) Cartera vencida detallada
# ----------------------------------------------------------------------------
def _nivel_alumno(matricula, grupo, niveles_roster):
    if matricula in niveles_roster and niveles_roster[matricula]:
        return niveles_roster[matricula].capitalize()
    g = (grupo or '').upper()
    if 'PREM' in g or 'PRE' in g:
        return 'Preescolar'
    if 'PRIM' in g:
        return 'Primaria'
    return 'Sin nivel'


@login_required
@requiere_membresia_activa_escuela
def reporte_cartera_detalle(request):
    empresa = _empresa_escuela(request)
    if not empresa:
        return redirect('dashboard_inicio')

    fechas = list(
        CarteraVencida.objects.filter(empresa=empresa)
        .values_list('fecha_corte', flat=True).distinct().order_by('-fecha_corte')
    )
    if not fechas:
        return render(request, 'escuelas/reporte_cartera_detalle.html', {'sin_datos': True})

    fecha_corte = _fecha_get(request, 'fecha_corte', fechas[0])
    if fecha_corte not in fechas:
        fecha_corte = fechas[0]
    anterior = next((f for f in fechas if f < fecha_corte), None)

    f_grupo = request.GET.get('grupo', '').strip()
    f_categoria = request.GET.get('categoria', '')
    f_bucket = request.GET.get('antiguedad', '')
    f_nivel = request.GET.get('nivel', '')
    f_q = request.GET.get('q', '').strip()
    try:
        f_min = Decimal(request.GET.get('min_monto') or '0')
    except Exception:  # noqa: BLE001
        f_min = CERO

    registros = list(CarteraVencida.objects.filter(empresa=empresa, fecha_corte=fecha_corte))
    niveles_roster = dict(Alumno.objects.filter(empresa=empresa).values_list('matricula', 'nivel'))

    alumnos = OrderedDict()
    for r in registros:
        meses = iu.meses_de_atraso(r.concepto, fecha_corte)
        b_clave, b_etq = iu.bucket_antiguedad(meses)
        nivel = _nivel_alumno(r.matricula, r.grupo, niveles_roster)
        if f_grupo and f_grupo.upper() not in (r.grupo or '').upper():
            continue
        if f_categoria and r.categoria != f_categoria:
            continue
        if f_bucket and b_clave != f_bucket:
            continue
        if f_nivel and nivel != f_nivel:
            continue
        if f_q and f_q.upper() not in f'{r.matricula} {r.nombre_alumno}'.upper():
            continue
        a = alumnos.setdefault(r.matricula, {
            'matricula': r.matricula, 'nombre': r.nombre_alumno, 'grupo': r.grupo or '',
            'nivel': nivel, 'total': CERO, 'conceptos': [], 'max_atraso': None,
        })
        a['total'] += r.monto
        a['conceptos'].append({
            'concepto': r.concepto, 'categoria': r.get_categoria_display(),
            'monto': r.monto, 'meses': meses, 'bucket': b_etq,
        })
        if meses is not None and (a['max_atraso'] is None or meses > a['max_atraso']):
            a['max_atraso'] = meses

    lista = [a for a in alumnos.values() if a['total'] >= f_min]
    lista.sort(key=lambda a: -a['total'])
    for a in lista:
        a['conceptos'].sort(key=lambda c: (-(c['meses'] if c['meses'] is not None else -99), c['concepto']))

    total = sum((a['total'] for a in lista), CERO)
    n_conceptos = sum(len(a['conceptos']) for a in lista)

    def agrupar(clave_fn, etiqueta_fn=lambda k: k):
        acc = defaultdict(lambda: {'total': CERO, 'alumnos': set(), 'n': 0})
        for a in lista:
            for c in a['conceptos']:
                k = clave_fn(a, c)
                acc[k]['total'] += c['monto']
                acc[k]['alumnos'].add(a['matricula'])
                acc[k]['n'] += 1
        out = [{'etiqueta': etiqueta_fn(k), 'total': v['total'], 'alumnos': len(v['alumnos']), 'n': v['n'],
                'pct': (v['total'] / total * 100) if total else CERO} for k, v in acc.items()]
        return sorted(out, key=lambda x: -x['total'])

    por_categoria = agrupar(lambda a, c: c['categoria'])
    orden_buckets = {etq: i for i, (_, etq, _) in enumerate(iu.BUCKETS_ANTIGUEDAD)}
    por_antiguedad = sorted(
        agrupar(lambda a, c: c['bucket']), key=lambda x: orden_buckets.get(x['etiqueta'], 99)
    )
    por_grupo = agrupar(lambda a, c: a['grupo'] or 'Sin grupo')
    por_grupo.sort(key=lambda x: x['etiqueta'])
    por_nivel = agrupar(lambda a, c: a['nivel'])

    # Variacion contra el corte anterior (sobre TODO el corte, sin filtros)
    comparativo = None
    if anterior:
        prev = defaultdict(lambda: CERO)
        for m, monto in CarteraVencida.objects.filter(empresa=empresa, fecha_corte=anterior).values_list('matricula', 'monto'):
            prev[m] += monto
        actual = defaultdict(lambda: CERO)
        for r in registros:
            actual[r.matricula] += r.monto
        nuevos = [m for m in actual if m not in prev]
        saldados = [m for m in prev if m not in actual]
        comparativo = {
            'fecha_anterior': anterior,
            'total_anterior': sum(prev.values(), CERO), 'total_actual': sum(actual.values(), CERO),
            'variacion': sum(actual.values(), CERO) - sum(prev.values(), CERO),
            'nuevos': len(nuevos), 'saldados': len(saldados),
            'monto_nuevos': sum((actual[m] for m in nuevos), CERO),
            'monto_saldados': sum((prev[m] for m in saldados), CERO),
        }

    if request.GET.get('exportar') == 'xlsx':
        wb = Workbook()
        ws = wb.active
        ws.title = 'Por alumno'
        _encabezado(ws, ['Matrícula', 'Alumno', 'Nivel', 'Grupo', '# conceptos', 'Máx. meses de atraso', 'Total adeudado'])
        for a in lista:
            ws.append([a['matricula'], a['nombre'], a['nivel'], a['grupo'], len(a['conceptos']),
                       a['max_atraso'] if a['max_atraso'] is not None else '', float(a['total'])])
        ws2 = wb.create_sheet('Detalle por concepto')
        _encabezado(ws2, ['Matrícula', 'Alumno', 'Nivel', 'Grupo', 'Concepto', 'Categoría', 'Meses de atraso', 'Antigüedad', 'Monto'])
        for a in lista:
            for c in a['conceptos']:
                ws2.append([a['matricula'], a['nombre'], a['nivel'], a['grupo'], c['concepto'], c['categoria'],
                            c['meses'] if c['meses'] is not None else '', c['bucket'], float(c['monto'])])
        ws3 = wb.create_sheet('Resumen')
        ws3.append([f'Cartera vencida al corte {fecha_corte:%d/%m/%Y}', float(total)])
        for titulo, bloque in (('Por categoría', por_categoria), ('Por antigüedad', por_antiguedad),
                               ('Por nivel', por_nivel), ('Por grupo', por_grupo)):
            ws3.append([])
            _encabezado_fila = [titulo, 'Alumnos', 'Conceptos', 'Total', '% del total']
            ws3.append(_encabezado_fila)
            for c in ws3[ws3.max_row]:
                c.font = Font(bold=True)
            for b in bloque:
                ws3.append([b['etiqueta'], b['alumnos'], b['n'], float(b['total']), float(b['pct']) / 100])
        return _excel_response(wb, f'cartera_vencida_detalle_{fecha_corte:%Y%m%d}.xlsx')

    return render(request, 'escuelas/reporte_cartera_detalle.html', {
        'fechas': fechas, 'fecha_corte': fecha_corte, 'alumnos': lista,
        'total': total, 'n_alumnos': len(lista), 'n_conceptos': n_conceptos,
        'promedio': (total / len(lista)) if lista else CERO,
        'por_categoria': por_categoria, 'por_antiguedad': por_antiguedad,
        'por_grupo': por_grupo, 'por_nivel': por_nivel, 'top10': lista[:10],
        'comparativo': comparativo,
        'categorias': CarteraVencida.CATEGORIA_CHOICES,
        'buckets': [(c, e) for c, e, _ in iu.BUCKETS_ANTIGUEDAD] + [('sin_mes', 'Sin mes en el concepto')],
        'niveles': ['Preescolar', 'Primaria', 'Sin nivel'],
        'f': {'grupo': f_grupo, 'categoria': f_categoria, 'antiguedad': f_bucket,
              'nivel': f_nivel, 'q': f_q, 'min_monto': request.GET.get('min_monto', '')},
        'labels_antig': [x['etiqueta'] for x in por_antiguedad],
        'datos_antig': [float(x['total']) for x in por_antiguedad],
    })
